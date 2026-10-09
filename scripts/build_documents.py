#!/usr/bin/env python3
"""Convert project Markdown manuscripts to formatted Word documents.
Basic Markdown subset: headings, tables, list items, paragraphs, emphasis markers.
"""
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'; OUT.mkdir(exist_ok=True)
FILES=[('docs/software/01_Software_Requirements_Specification.md','01_Software_Requirements_Specification.docx',False),
 ('docs/software/02_Architecture_Design.md','02_Software_Design_and_Architecture.docx',False),
 ('docs/software/03_Database_Design.md','03_Database_Design_and_Data_Dictionary.docx',False),
 ('docs/software/04_Test_and_QA_Plan.md','04_Testing_and_QA_Plan.docx',False),
 ('docs/software/05_Deployment_and_Operations_Manual.md','05_Deployment_and_Operations_Manual.docx',False),
 ('docs/research/proposal.md','06_Research_Proposal_Chapters_1_to_3.docx',True),
 ('docs/research/full_project.md','07_Research_Project_Chapters_1_to_5.docx',True)]

def shade(cell,color):
 tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd');sh.set(qn('w:fill'),color);tcPr.append(sh)

def keep_with_next(p):p.paragraph_format.keep_with_next=True

def field(p):
 run=p.add_run(); a=OxmlElement('w:fldChar');a.set(qn('w:fldCharType'),'begin'); run._r.append(a)
 b=OxmlElement('w:instrText');b.set(qn('xml:space'),'preserve');b.text=' PAGE ';run._r.append(b)
 c=OxmlElement('w:fldChar');c.set(qn('w:fldCharType'),'end');run._r.append(c)

def add_rich(p,s):
 s=s.replace('`','');pieces=re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)',s)
 for piece in pieces:
  if piece.startswith('**') and piece.endswith('**'):p.add_run(piece[2:-2]).bold=True
  elif piece.startswith('*') and piece.endswith('*'):p.add_run(piece[1:-1]).italic=True
  elif piece:p.add_run(piece)

def make(infile,outfile,research):
 lines=(ROOT/infile).read_text().splitlines();doc=Document();sec=doc.sections[0];sec.page_height=Inches(11.69);sec.page_width=Inches(8.27)
 sec.left_margin=Inches(1.15);sec.right_margin=Inches(.9);sec.top_margin=Inches(.94);sec.bottom_margin=Inches(.83)
 sec.header_distance=Inches(.36);sec.footer_distance=Inches(.38)
 normal=doc.styles['Normal'];normal.font.name='Times New Roman' if research else 'Aptos';normal.font.size=Pt(12 if research else 10.5);normal.font.color.rgb=RGBColor(38,50,65)
 normal.paragraph_format.line_spacing=1.5 if research else 1.21;normal.paragraph_format.space_after=Pt(7)
 for level,size in [(1,15.5),(2,13),(3,11.5)]:
  st=doc.styles[f'Heading {level}'];st.font.name=normal.font.name;st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(12,64,80) if not research else RGBColor(26,37,49);st.paragraph_format.space_before=Pt(18 if level==1 else 12);st.paragraph_format.space_after=Pt(7);st.paragraph_format.keep_with_next=True
 header=sec.header.paragraphs[0];header.text='RSCDS  |  '+('ACADEMIC RESEARCH' if research else 'ENGINEERING DOCUMENTATION');header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 for rr in header.runs:rr.font.size=Pt(8);rr.font.color.rgb=RGBColor(112,130,140)
 ft=sec.footer.paragraphs[0];ft.alignment=WD_ALIGN_PARAGRAPH.CENTER;ft.add_run('RSCDS • October 2026  |  Page ');field(ft)
 first_heading=next((x[2:].strip() for x in lines if x.startswith('# ')),'RSCDS')
 title_lines=[];body_start=1
 for i,line in enumerate(lines[1:],start=1):
  clean=line.strip()
  if not clean:continue
  if clean.startswith('## ') and not title_lines:
   title_lines.append(clean[3:]);continue
  if clean.startswith('**') and len(title_lines)<9:
   title_lines.append(clean.strip('*'));continue
  body_start=i;break
 cover=doc.add_paragraph();cover.alignment=WD_ALIGN_PARAGRAPH.CENTER;cover.paragraph_format.space_before=Pt(80)
 r=cover.add_run(first_heading);r.font.bold=True;r.font.name=normal.font.name;r.font.size=Pt(18);r.font.color.rgb=RGBColor(13,69,80)
 for value in title_lines[:7]:
  pp=doc.add_paragraph();pp.alignment=WD_ALIGN_PARAGRAPH.CENTER;pp.paragraph_format.space_before=Pt(10);rr=pp.add_run(re.sub(r'\*\*','',value));rr.font.size=Pt(12.5 if research else 11.2)
 note=doc.add_paragraph();note.alignment=WD_ALIGN_PARAGRAPH.CENTER;note.paragraph_format.space_before=Pt(25)
 note.add_run('RESEARCH DRAFT • INTEGRATION/EMPIRICAL VALIDATION STATUS EXPLICITLY REPORTED' if research else 'VERSION 1.0 • IMPLEMENTATION BASELINE').bold=True
 doc.add_page_break()
 reading_refs=False
 i=body_start
 while i<len(lines):
  s=lines[i].strip()
  if not s:i+=1;continue
  if research and s.startswith('# ') and ('CHAPTER ' in s or s in ('# ABSTRACT','# REFERENCES')):
   if s!='# ABSTRACT':doc.add_page_break()
   heading=doc.add_heading(s[2:],1);heading.alignment=WD_ALIGN_PARAGRAPH.CENTER;reading_refs=(s=='# REFERENCES');i+=1;continue
  if s.startswith('# '):
   if not research:doc.add_heading(s[2:],1)
   i+=1;continue
  if s.startswith('## '):
   if research and ('Satellite Image-Based' in s):i+=1;continue
   doc.add_heading(s[3:],2);i+=1;continue
  if s.startswith('### '):doc.add_heading(s[4:],3);i+=1;continue
  if s.startswith('|') and i+1<len(lines) and lines[i+1].strip().startswith('|---'):
   rows=[];j=i
   while j<len(lines) and lines[j].strip().startswith('|'):
    if j!=i+1:rows.append([q.strip() for q in lines[j].strip().strip('|').split('|')])
    j+=1
   if rows:
    count=max(map(len,rows));tbl=doc.add_table(rows=0,cols=count);tbl.style='Table Grid';tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    for rid,vals in enumerate(rows):
     cells=tbl.add_row().cells
     for cid,txt in enumerate(vals):
      pp=cells[cid].paragraphs[0];add_rich(pp,txt);pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1.0
      for run in pp.runs:run.font.size=Pt(8 if count>3 else 9)
      cells[cid].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
      if rid==0:shade(cells[cid],'DCEEF0');
      if rid==0:
       for run in pp.runs:run.bold=True
    for row in tbl.rows:
     for cell in row.cells:
      tcMar=cell._tc.get_or_add_tcPr();
      # no forced row heights: allow wrapping naturally
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2)
   i=j;continue
  if s.startswith('- '):
   pp=doc.add_paragraph(style='List Bullet');add_rich(pp,s[2:]);pp.paragraph_format.left_indent=Inches(.25);i+=1;continue
  if re.match(r'^\d+[.)] ',s):
   pp=doc.add_paragraph(style='List Number');add_rich(pp,re.sub(r'^\d+[.)] ','',s));i+=1;continue
  if s.startswith('**') and len(s)<150 and not s.endswith('.') and i<15:
   i+=1;continue
  pp=doc.add_paragraph();add_rich(pp,s);pp.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
  if reading_refs:
   pp.paragraph_format.first_line_indent=Inches(-.25);pp.paragraph_format.left_indent=Inches(.25)
  elif research:pp.paragraph_format.first_line_indent=Inches(.27)
  i+=1
 path=OUT/outfile;doc.save(path);return path

if __name__=='__main__':
 for src,dst,research in FILES:
  p=make(src,dst,research);print(p,p.stat().st_size)
